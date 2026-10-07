import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10)
GREEN = Font(name=F, size=10, color="008000")
BOLD = Font(name=F, size=10, bold=True)
HFONT = Font(name=F, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=F, size=14, bold=True)
HFILL = PatternFill("solid", fgColor="1F3A5F")
CHFILL = PatternFill("solid", fgColor="D9E2F3")
YEL = PatternFill("solid", fgColor="FFFF00")
TOTFILL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
N2 = '#,##0.00;-#,##0.00;"-"'
N0 = '#,##0;-#,##0;"-"'
N3 = '#,##0.000;-#,##0.000;"-"'

wb = Workbook()


def hdr(ws, row, labels, widths=None):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=l)
        c.font, c.fill, c.alignment, c.border = HFONT, HFILL, CENTER, BRD
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, ref, val, font=BLACK, fmt=None, fill=None, note=None):
    c = ws[ref]
    c.value = val
    c.font = font
    c.border = BRD
    c.alignment = WRAP
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, "Cantidades")
    return c


# ---------------------------------------------------------------- SUPUESTOS
S = wb.active
S.title = "Supuestos"
S["A1"] = "SUPUESTOS Y PARÁMETROS DE CÁLCULO"
S["A1"].font = TITLE
S["A2"] = ("Celdas en azul con fondo amarillo = parámetros editables. Al cambiarlos se recalculan "
           "todas las cantidades derivadas. Fuente indicada en la columna D.")
S["A2"].font = Font(name=F, size=9, italic=True)
hdr(S, 3, ["Parámetro", "Valor", "Und", "Fuente / criterio"], [52, 12, 10, 80])
SUP = [
    ("peso3", "Peso nominal barra #3 (3/8\")", 0.560, "kg/m", "NSR-10 / NTC 2289"),
    ("peso4", "Peso nominal barra #4 (1/2\")", 0.994, "kg/m", "NSR-10 / NTC 2289"),
    ("desp_acero", "Desperdicio acero de refuerzo", 0.05, "%", "Criterio usual de obra (traslapos y cortes)"),
    ("desp_conc", "Desperdicio concreto", 0.05, "%", "Criterio usual de obra"),
    ("und_m2", "Unidades de mampostería por m² (bloque arcilla 10×20×30 cm, junta 1 cm)", 15.4, "un/m²",
     "1/(0.31×0.21). AJUSTAR si la unidad especificada es otra (p.ej. tolete 6×12×24 ≈ 54 un/m² en soga)"),
    ("desp_mamp", "Desperdicio unidades de mampostería", 0.05, "%", "Criterio usual de obra"),
    ("mort_pega", "Mortero de pega por m² de muro (bloque e=10 cm)", 0.012, "m³/m²", "Rendimiento de referencia; ajustar según unidad y espesor de junta"),
    ("caras", "Caras de pañete/estuco/pintura por muro", 2, "caras", "Supuesto: todos los muros se revocan por ambas caras (bruto, sin descontar enchapes)"),
    ("e_alist", "Espesor de alistado (mortero de nivelación) de piso", 0.04, "m", "Supuesto de diseño; confirmar con nivel de acabado"),
    ("sobreancho", "Sobreancho de excavación por lado (foso)", 0.30, "m", "Espacio de trabajo para formaleta exterior"),
    ("prof_exc", "Profundidad de excavación foso (desde N-0.10)", 1.45, "m", "N-0.10 a N-1.50 (EST-02) + 0.05 de solado"),
    ("expansion", "Factor de expansión del suelo (sobrantes)", 1.30, "", "Criterio usual para material común"),
    ("e_solado", "Espesor solado de limpieza", 0.05, "m", "Recomendado bajo losa en contacto con suelo (no dibujado en EST-02)"),
    ("sob_solado", "Sobreancho solado por lado", 0.10, "m", "Criterio usual"),
    ("dens_acero", "Densidad acero estructural", 7850, "kg/m³", "Valor estándar"),
    ("he200", "Peso perfil HE 200 (se asume HEB 200)", 61.3, "kg/m", "HEB 200 = 61.3 kg/m; si es HEA 200 usar 42.3 kg/m. EST-02 no lo precisa"),
    ("apoyo_dintel", "Apoyo de dintel por lado", 0.20, "m", "Criterio usual en mampostería"),
    ("e_muro_foso_imp", "Recubrimiento de impermeabilización interior foso", 1, "", "1 = incluir (recomendado para foso bajo nivel de piso)"),
]
SREF = {}
r = 4
for key, name, val, und, src in SUP:
    put(S, f"A{r}", name)
    fmt = "0%" if und == "%" else ("#,##0.000" if isinstance(val, float) and val < 1 else "#,##0.00")
    put(S, f"B{r}", val, BLUE, fmt, YEL)
    put(S, f"C{r}", und if und != "%" else "")
    put(S, f"D{r}", src)
    SREF[key] = f"Supuestos!$B${r}"
    r += 1

# ---------------------------------------------------------------- DATOS REVIT
D = wb.create_sheet("Datos Revit")
D["A1"] = "DATOS DE ENTRADA — EXTRAÍDOS DEL MODELO REVIT (MULTIF H - RVT.rvt)"
D["A1"].font = TITLE
D["A2"] = ("Fuente: dashboard index.html (consulta directa al modelo). Solo niveles P1–P5 terminados; "
           "P6/P7 y cubierta en diseño, NO incluidos.")
D["A2"].font = Font(name=F, size=9, italic=True)
for col, w in zip("ABCDEFGHIJ", [44, 10, 10, 10, 10, 10, 10, 10, 12, 40]):
    D.column_dimensions[col].width = w

NIV = ["P1", "P2", "P3", "P4", "P5"]
DR = {}

# Superficies y muros
r = 4
D[f"A{r}"] = "1. SUPERFICIES Y MUROS POR NIVEL"; D[f"A{r}"].font = BOLD
r += 1
hdr(D, r, ["Concepto", "Und"] + NIV + ["Total", "", "Nota"])
rows = [
    ("area_cub", "Área cubierta", "m²", [128.03, 110.53, 109.91, 110.04, 108.30], "Programa de superficies Revit"),
    ("area_ext", "Área exterior", "m²", [37.09, 0, 0, 0, 0], "Solo P1"),
    ("n_muros", "N° de muros de mampostería", "un", [24, 29, 26, 26, 26], ""),
    ("a_muros", "Área de muros de mampostería (neta, una cara)", "m²", [339.40, 251.10, 256.42, 256.08, 256.15],
     "Revit descuenta vanos de puertas/ventanas. Sin guardaescoba"),
]
for key, name, und, vals, nota in rows:
    r += 1
    put(D, f"A{r}", name); put(D, f"B{r}", und)
    for i, v in enumerate(vals):
        put(D, f"{get_column_letter(3+i)}{r}", v, BLUE, N2 if und == "m²" else N0)
    put(D, f"H{r}", f"=SUM(C{r}:G{r})", BOLD, N2 if und == "m²" else N0)
    put(D, f"J{r}", nota)
    DR[key] = r

# Puertas
r += 2
D[f"A{r}"] = "2. PUERTAS — por función y medida"; D[f"A{r}"].font = BOLD
r += 1
hdr(D, r, ["Función", "Ancho (m)", "Alto (m)"] + NIV + ["Total", "Nota"])
PUERTAS = [
    ("Puerta de garaje", 2.40, 2.10, [1, 0, 0, 0, 0], ""),
    ("Puerta principal de acceso", 0.90, 2.10, [1, 0, 0, 0, 0], ""),
    ("Puerta de acceso secundario", 0.915, 2.134, [1, 0, 0, 0, 0], "Medida imperial 3'×7'"),
    ("Puerta corrediza social (balcón)", 1.20, 2.10, [1, 0, 0, 0, 0], ""),
    ("Puerta corrediza de habitación/closet", 0.90, 2.10, [3, 0, 0, 0, 0], ""),
    ("Puerta corrediza de baño/closet menor", 0.70, 2.10, [0, 3, 2, 2, 2], ""),
    ("Puerta batiente de habitación", 0.80, 2.10, [2, 1, 0, 0, 0], ""),
    ("Puerta batiente de habitación", 0.90, 2.10, [2, 0, 3, 3, 3], ""),
    ("Puerta batiente de baño/closet", 0.70, 2.10, [0, 1, 1, 1, 1], ""),
    ("Puerta principal de apartamento", 1.00, 2.10, [0, 1, 0, 0, 0], ""),
    ("Puerta-ventanal corrediza (balcón)", 1.60, 2.30, [0, 0, 1, 1, 1], ""),
    ("Puerta corrediza de vestier (herraje Häfele)", None, None, [0, 1, 1, 1, 1], "Revit reporta ≈7.73 m² de hoja; sin ancho/alto definido"),
]
DR["puertas"] = []
for name, w, h, vals, nota in PUERTAS:
    r += 1
    put(D, f"A{r}", name)
    put(D, f"B{r}", w, BLUE, "0.000")
    put(D, f"C{r}", h, BLUE, "0.000")
    for i, v in enumerate(vals):
        put(D, f"{get_column_letter(4+i)}{r}", v, BLUE, N0)
    put(D, f"I{r}", f"=SUM(D{r}:H{r})", BOLD, N0)
    put(D, f"J{r}", nota)
    DR["puertas"].append(r)
r += 1
p0, p1 = DR["puertas"][0], DR["puertas"][-1]
put(D, f"A{r}", "TOTAL PUERTAS", BOLD, fill=TOTFILL)
for c in "DEFGHI":
    put(D, f"{c}{r}", f"=SUM({c}{p0}:{c}{p1})", BOLD, N0, TOTFILL)
DR["puertas_tot"] = r

# Ventanas
r += 2
D[f"A{r}"] = "3. VENTANAS Y VENTANALES — por función y medida"; D[f"A{r}"].font = BOLD
r += 1
hdr(D, r, ["Función", "Ancho (m)", "Alto (m)"] + NIV + ["Total", "Nota"])
VENT = [
    ("Ventana de habitación", 0.80, 0.80, [0, 3, 8, 8, 8], ""),
    ("Ventana de baño/cocina", 0.91, 0.91, [0, 2, 1, 1, 1], ""),
    ("Ventanal de circulación (hall/escalera)", 0.90, 2.41, [0, 1, 1, 1, 1], "Panel de muro cortina"),
    ("Ventanal de fachada — acceso", 0.68, 2.20, [2, 0, 0, 0, 0], "Panel de muro cortina"),
]
DR["vent"] = []
for name, w, h, vals, nota in VENT:
    r += 1
    put(D, f"A{r}", name)
    put(D, f"B{r}", w, BLUE, "0.000"); put(D, f"C{r}", h, BLUE, "0.000")
    for i, v in enumerate(vals):
        put(D, f"{get_column_letter(4+i)}{r}", v, BLUE, N0)
    put(D, f"I{r}", f"=SUM(D{r}:H{r})", BOLD, N0)
    put(D, f"J{r}", nota)
    DR["vent"].append(r)
r += 1
v0, v1 = DR["vent"][0], DR["vent"][-1]
put(D, f"A{r}", "TOTAL VENTANAS", BOLD, fill=TOTFILL)
for c in "DEFGHI":
    put(D, f"{c}{r}", f"=SUM({c}{v0}:{c}{v1})", BOLD, N0, TOTFILL)

# Sanitarios
r += 2
D[f"A{r}"] = "4. APARATOS SANITARIOS — verificados elemento por elemento"; D[f"A{r}"].font = BOLD
r += 1
hdr(D, r, ["Aparato", "", ""] + NIV + ["Total", "Nota"])
SAN = [
    ("san", "Sanitario", [2, 3, 3, 3, 3]),
    ("cab", "Cabina de ducha", [0, 3, 3, 3, 3]),
    ("reg", "Regadera de ducha", [0, 3, 3, 3, 3]),
    ("dacc", "Ducha accesible (con barras)", [1, 0, 0, 0, 0]),
    ("lav", "Lavamanos de baño", [2, 3, 3, 3, 3]),
    ("mamp", "Mampara de baño", [0, 0, 3, 3, 3]),
]
for key, name, vals in SAN:
    r += 1
    put(D, f"A{r}", name)
    for i, v in enumerate(vals):
        put(D, f"{get_column_letter(4+i)}{r}", v, BLUE, N0)
    put(D, f"I{r}", f"=SUM(D{r}:H{r})", BOLD, N0)
    DR["s_" + key] = r
r += 1
s0, s1 = DR["s_san"], DR["s_mamp"]
put(D, f"A{r}", "TOTAL APARATOS", BOLD, fill=TOTFILL)
for c in "DEFGHI":
    put(D, f"{c}{r}", f"=SUM({c}{s0}:{c}{s1})", BOLD, N0, TOTFILL)

# ---------------------------------------------------------------- FOSO
FO = wb.create_sheet("Foso ascensor")
FO["A1"] = "MEMORIA DE CÁLCULO — FOSO DE ASCENSOR (plano EST-02, versión 01, 19/06/2026)"
FO["A1"].font = TITLE
FO["A2"] = ("Diseño estructural: A3 Ingeniería. Escala 1:25. Cotas en m. Valores en azul tomados del plano; "
            "las cantidades son fórmulas.")
FO["A2"].font = Font(name=F, size=9, italic=True)
for col, w in zip("ABCDEFGHIJKL", [44, 9, 10, 12, 12, 11, 12, 10, 12, 10, 12, 44]):
    FO.column_dimensions[col].width = w

r = 4
FO[f"A{r}"] = "A. GEOMETRÍA (según planta, cortes A-A y B-B)"; FO[f"A{r}"].font = BOLD
r += 1
hdr(FO, r, ["Dato", "Valor", "Und", "Fuente"])
GEO = [
    ("intX", "Luz libre interior del foso — sentido X", 1.50, "Planta foso"),
    ("intY", "Luz libre interior del foso — sentido Y", 1.90, "Planta foso"),
    ("e", "Espesor muros de concreto", 0.15, "Planta foso: Muro e=0.15m"),
    ("col", "Lado columna cuadrada (4 und, en esquinas)", 0.40, "Columna 40×40"),
    ("ncol", "Número de columnas", 4, "Planta foso"),
    ("luzX", "Distancia libre entre caras de columnas — X", 1.30, "Planta foso"),
    ("luzY", "Distancia libre entre caras de columnas — Y", 1.70, "Planta foso"),
    ("H", "Altura total foso (N-0.10 a N-1.50)", 1.40, "Cortes A-A / B-B"),
    ("eL", "Espesor losa de fondo", 0.20, "Losa e=0.20m"),
]
G = {}
for key, name, val, src in GEO:
    r += 1
    put(FO, f"A{r}", name)
    put(FO, f"B{r}", val, BLUE, "0.00" if key != "ncol" else "0")
    put(FO, f"C{r}", "un" if key == "ncol" else "m")
    put(FO, f"D{r}", src)
    G[key] = f"$B${r}"
r += 1
put(FO, f"A{r}", "Losa — dimensión exterior X (luz + 2 muros)"); put(FO, f"B{r}", f"={G['intX']}+2*{G['e']}", fmt="0.00"); put(FO, f"C{r}", "m"); G["LX"] = f"$B${r}"
r += 1
put(FO, f"A{r}", "Losa — dimensión exterior Y (luz + 2 muros)"); put(FO, f"B{r}", f"={G['intY']}+2*{G['e']}", fmt="0.00"); put(FO, f"C{r}", "m"); G["LY"] = f"$B${r}"
r += 1
put(FO, f"A{r}", "Huella total con columnas — X"); put(FO, f"B{r}", f"={G['luzX']}+2*{G['col']}", fmt="0.00"); put(FO, f"C{r}", "m"); G["TX"] = f"$B${r}"
r += 1
put(FO, f"A{r}", "Huella total con columnas — Y"); put(FO, f"B{r}", f"={G['luzY']}+2*{G['col']}", fmt="0.00"); put(FO, f"C{r}", "m"); G["TY"] = f"$B${r}"
r += 1
put(FO, f"A{r}", "Traslapo columna–losa por esquina (lado)"); put(FO, f"B{r}", f"={G['col']}-({G['TX']}-{G['LX']})/2", fmt="0.000"); put(FO, f"C{r}", "m"); G["ov"] = f"$B${r}"
r += 1
put(FO, f"A{r}", "Altura libre muros (sobre losa)"); put(FO, f"B{r}", f"={G['H']}-{G['eL']}", fmt="0.00"); put(FO, f"C{r}", "m"); G["hm"] = f"$B${r}"
r += 1
put(FO, f"A{r}", "Profundidad útil del foso (N-0.10 a cara superior losa)"); put(FO, f"B{r}", f"={G['hm']}", fmt="0.00"); put(FO, f"C{r}", "m")
put(FO, f"D{r}", "Verificar contra la ficha técnica del ascensor elegido")

# Concreto
r += 2
FO[f"A{r}"] = "B. CONCRETO"; FO[f"A{r}"].font = BOLD
r += 1
hdr(FO, r, ["Elemento", "Cant.", "Und", "Cálculo"])
FC = {}
r += 1
put(FO, f"A{r}", "Columnas 40×40 (altura total foso)")
put(FO, f"B{r}", f"={G['ncol']}*{G['col']}^2*{G['H']}", fmt=N3); put(FO, f"C{r}", "m³")
put(FO, f"D{r}", "4 × 0.40 × 0.40 × 1.40"); FC["col"] = r
r += 1
put(FO, f"A{r}", "Losa de fondo e=0.20 (descontando columnas)")
put(FO, f"B{r}", f"={G['LX']}*{G['LY']}*{G['eL']}-{G['ncol']}*{G['ov']}^2*{G['eL']}", fmt=N3); put(FO, f"C{r}", "m³")
put(FO, f"D{r}", "1.80 × 2.20 × 0.20 − 4 × 0.25² × 0.20"); FC["losa"] = r
r += 1
put(FO, f"A{r}", "Muros e=0.15 (entre columnas, sobre losa)")
put(FO, f"B{r}", f"=(2*{G['luzX']}+2*{G['luzY']})*{G['hm']}*{G['e']}", fmt=N3); put(FO, f"C{r}", "m³")
put(FO, f"D{r}", "(2×1.30 + 2×1.70) × 1.20 × 0.15"); FC["muro"] = r
r += 1
put(FO, f"A{r}", "Solado de limpieza (concreto pobre)")
put(FO, f"B{r}", f"=({G['LX']}+2*{SREF['sob_solado']})*({G['LY']}+2*{SREF['sob_solado']})*{SREF['e_solado']}", fmt=N3)
put(FO, f"C{r}", "m³"); put(FO, f"D{r}", "(1.80+0.20) × (2.20+0.20) × 0.05 — recomendado"); FC["solado"] = r
r += 1
put(FO, f"A{r}", "TOTAL CONCRETO ESTRUCTURAL (sin desperdicio)", BOLD, fill=TOTFILL)
put(FO, f"B{r}", f"=B{FC['col']}+B{FC['losa']}+B{FC['muro']}", BOLD, N3, TOTFILL); put(FO, f"C{r}", "m³", fill=TOTFILL)
FC["tot"] = r

# Movimiento de tierra
r += 2
FO[f"A{r}"] = "C. MOVIMIENTO DE TIERRA Y FORMALETA"; FO[f"A{r}"].font = BOLD
r += 1
hdr(FO, r, ["Concepto", "Cant.", "Und", "Cálculo"])
r += 1
put(FO, f"A{r}", "Excavación manual")
put(FO, f"B{r}", f"=({G['TX']}+2*{SREF['sobreancho']})*({G['TY']}+2*{SREF['sobreancho']})*{SREF['prof_exc']}", fmt=N2)
put(FO, f"C{r}", "m³"); put(FO, f"D{r}", "(2.10+0.60) × (2.50+0.60) × 1.45"); FC["exc"] = r
r += 1
put(FO, f"A{r}", "Volumen ocupado por el foso + solado")
put(FO, f"B{r}", f"=({G['LX']}*{G['LY']}+{G['ncol']}*({G['col']}^2-{G['ov']}^2))*{G['H']}+B{FC['solado']}", fmt=N2)
put(FO, f"C{r}", "m³"); put(FO, f"D{r}", "Huella real (losa + salientes de columnas) × 1.40 + solado"); FC["ocup"] = r
r += 1
put(FO, f"A{r}", "Relleno compactado perimetral (material seleccionado)")
put(FO, f"B{r}", f"=B{FC['exc']}-B{FC['ocup']}", fmt=N2); put(FO, f"C{r}", "m³")
put(FO, f"D{r}", "Excavación − volumen ocupado. Compactar en capas ≤ 0.20 m"); FC["rell"] = r
r += 1
put(FO, f"A{r}", "Retiro de sobrantes (suelto)")
put(FO, f"B{r}", f"=(B{FC['exc']}-B{FC['rell']})*{SREF['expansion']}", fmt=N2); put(FO, f"C{r}", "m³")
put(FO, f"D{r}", "(Excavación − relleno) × factor de expansión"); FC["ret"] = r
r += 1
put(FO, f"A{r}", "Formaleta muros — cara interior")
put(FO, f"B{r}", f"=2*({G['intX']}+{G['intY']})*{G['hm']}", fmt=N2); put(FO, f"C{r}", "m²")
put(FO, f"D{r}", "Perímetro interior 6.80 × 1.20"); FC["f1"] = r
r += 1
put(FO, f"A{r}", "Formaleta muros — cara exterior")
put(FO, f"B{r}", f"=(2*{G['luzX']}+2*{G['luzY']})*{G['H']}", fmt=N2); put(FO, f"C{r}", "m²")
put(FO, f"D{r}", "Tramos entre columnas 6.00 × 1.40"); FC["f2"] = r
r += 1
put(FO, f"A{r}", "Formaleta columnas — caras exteriores expuestas")
put(FO, f"B{r}", f"={G['ncol']}*(2*{G['col']}+2*({G['col']}-{G['ov']}))*{G['H']}", fmt=N2); put(FO, f"C{r}", "m²")
put(FO, f"D{r}", "4 × (2×0.40 + 2×0.15) × 1.40"); FC["f3"] = r
r += 1
put(FO, f"A{r}", "TOTAL FORMALETA", BOLD, fill=TOTFILL)
put(FO, f"B{r}", f"=SUM(B{FC['f1']}:B{FC['f3']})", BOLD, N2, TOTFILL); put(FO, f"C{r}", "m²", fill=TOTFILL)
FC["form"] = r
r += 1
put(FO, f"A{r}", "Impermeabilización interior (fondo + muros)")
put(FO, f"B{r}", f"=({G['intX']}*{G['intY']}+2*({G['intX']}+{G['intY']})*{G['hm']})*{SREF['e_muro_foso_imp']}", fmt=N2)
put(FO, f"C{r}", "m²"); put(FO, f"D{r}", "1.50×1.90 + 6.80×1.20 — no especificado en EST-02, recomendado"); FC["imp"] = r

# Despiece
r += 2
FO[f"A{r}"] = "D. CARTILLA DE DESPIECE — ACERO DE REFUERZO fy = 420 MPa"; FO[f"A{r}"].font = BOLD
r += 1
hdr(FO, r, ["Elemento / posición", "Barra", "Separación (m)", "Long. a distribuir (m)",
            "N° barras por capa/elem.", "Capas o elementos", "N° barras total", "L barra (m)",
            "Long. total (m)", "kg/m", "Peso (kg)", "Fuente en plano"])
DESP = [
    ("Losa fondo — barras sentido X (corto)", "#4", 0.20, 2.08, None, 2, 2.20, "Corte A-A: #4 c/0.20 L=2.20 (doble parrilla)"),
    ("Losa fondo — barras sentido Y (largo)", "#4", 0.20, 1.68, None, 2, 2.60, "Corte B-B: #4 c/0.20 L=2.60 (doble parrilla)"),
    ("Muros sentido X — horizontales", "#4", 0.20, f"={G['hm']}", None, 2, 2.10, "Planta refuerzo: #4 c/0.20 L=2.10"),
    ("Muros sentido Y — horizontales", "#4", 0.20, f"={G['hm']}", None, 2, 2.50, "Planta refuerzo: #4 c/0.20 L=2.50"),
    ("Muros sentido X — verticales (gancho a losa)", "#4", 0.20, f"={G['luzX']}", None, 2, 1.75, "Cortes A-A/B-B: #4 c/0.20 L=1.75"),
    ("Muros sentido Y — verticales (gancho a losa)", "#4", 0.20, f"={G['luzY']}", None, 2, 1.75, "Cortes A-A/B-B: #4 c/0.20 L=1.75"),
    ("Columnas 40×40 — longitudinal", "#4", None, None, 4, 4, 1.85, "Confirmado: 4 #4 por columna (L=1.85 del detalle)"),
    ("Columnas 40×40 — flejes cerrados", "#3", 0.075, f"={G['H']}", None, 4, 1.48, "Confirmado: fleje #3 @0.075 L=1.48"),
]
d0 = r + 1
for name, bar, sep, dist, nfix, mult, L, src in DESP:
    r += 1
    put(FO, f"A{r}", name); put(FO, f"B{r}", bar, fmt=None)
    put(FO, f"C{r}", sep, BLUE if sep else BLACK, "0.000")
    put(FO, f"D{r}", dist, BLUE if isinstance(dist, float) else BLACK, "0.00")
    if nfix is not None:
        put(FO, f"E{r}", nfix, BLUE, N0)
    else:
        put(FO, f"E{r}", f"=ROUNDUP(D{r}/C{r},0)+1", fmt=N0)
    put(FO, f"F{r}", mult, BLUE, N0)
    put(FO, f"G{r}", f"=E{r}*F{r}", fmt=N0)
    put(FO, f"H{r}", L, BLUE, "0.00")
    put(FO, f"I{r}", f"=G{r}*H{r}", fmt=N2)
    put(FO, f"J{r}", f"={SREF['peso4']}" if bar == "#4" else f"={SREF['peso3']}", GREEN, "0.000")
    put(FO, f"K{r}", f"=I{r}*J{r}", fmt=N2)
    put(FO, f"L{r}", src)
d1 = r
r += 1
put(FO, f"A{r}", "Subtotal #4 (sin desperdicio)", BOLD, fill=TOTFILL)
put(FO, f"I{r}", f'=SUMIF(B{d0}:B{d1},"#4",I{d0}:I{d1})', BOLD, N2, TOTFILL)
put(FO, f"K{r}", f'=SUMIF(B{d0}:B{d1},"#4",K{d0}:K{d1})', BOLD, N2, TOTFILL); FC["k4"] = r
r += 1
put(FO, f"A{r}", "Subtotal #3 (sin desperdicio)", BOLD, fill=TOTFILL)
put(FO, f"I{r}", f'=SUMIF(B{d0}:B{d1},"#3",I{d0}:I{d1})', BOLD, N2, TOTFILL)
put(FO, f"K{r}", f'=SUMIF(B{d0}:B{d1},"#3",K{d0}:K{d1})', BOLD, N2, TOTFILL); FC["k3"] = r
r += 1
put(FO, f"A{r}", "Cuantía de referencia (kg acero / m³ concreto)", BOLD)
put(FO, f"K{r}", f"=(K{FC['k4']}+K{FC['k3']})/B{FC['tot']}", BOLD, "0.0")

# Estructura metálica
r += 2
FO[f"A{r}"] = "E. ESTRUCTURA METÁLICA — DETALLE CONEXIÓN 1 (perfil HE 200 a columna 40×40)"; FO[f"A{r}"].font = BOLD
r += 1
hdr(FO, r, ["Elemento", "Cant.", "Und", "Especificación / cálculo"])
r += 1
put(FO, f"A{r}", "Platina base A36 300×300×16 mm — unidades"); put(FO, f"B{r}", f"={G['ncol']}", fmt=N0); put(FO, f"C{r}", "un")
put(FO, f"D{r}", "1 por columna"); FC["pl_un"] = r
r += 1
put(FO, f"A{r}", "Platina base A36 300×300×16 mm — peso"); put(FO, f"B{r}", f"=B{FC['pl_un']}*0.30*0.30*0.016*{SREF['dens_acero']}", fmt=N2)
put(FO, f"C{r}", "kg"); put(FO, f"D{r}", "0.30 × 0.30 × 0.016 × 7850 por platina"); FC["pl_kg"] = r
r += 1
put(FO, f"A{r}", "Pernos de anclaje Ø5/8\" con gancho"); put(FO, f"B{r}", f"=4*{G['ncol']}", fmt=N0); put(FO, f"C{r}", "un")
put(FO, f"D{r}", "4 por platina. Longitud ≈ 0.34 m + gancho 0.10 m (detalle perno); embebido 314 mm"); FC["pernos"] = r
r += 1
put(FO, f"A{r}", "Soldadura filete 5/16\" alrededor del perfil"); put(FO, f"B{r}", f"={G['ncol']}*1.15", fmt=N2); put(FO, f"C{r}", "ml")
put(FO, f"D{r}", "≈1.15 m de cordón por perfil HE 200 (contorno del perfil)"); FC["sold"] = r
r += 1
put(FO, f"A{r}", "Perfil HE 200 — columnas del ducto (arranque)"); put(FO, f"B{r}", f"={G['ncol']}", fmt=N0); put(FO, f"C{r}", "un")
put(FO, f"D{r}", "Longitud = altura del ducto hasta sala de máquinas/cubierta: NO definida en EST-02"); FC["he"] = r

# ---------------------------------------------------------------- CANTIDADES
C = wb.create_sheet("Cantidades", 0)
C["A1"] = "CUADRO DE CANTIDADES DE OBRA — MULTIFAMILIAR HERRERA (Puerto Berrío, Antioquia)"
C["A1"].font = TITLE
C["A2"] = ("Alcance: niveles P1–P5 terminados en el modelo Revit + foso de ascensor (plano EST-02). "
           "Sin precios: el presupuesto se maneja aparte. Fuente: R = medido en Revit · E = plano estructural EST-02 · "
           "D = derivada por fórmula (ver hoja Supuestos) · P = por confirmar.")
C["A2"].font = Font(name=F, size=9, italic=True)
C.merge_cells("A2:L2")
C.row_dimensions[2].height = 30
hdr(C, 4, ["Ítem", "Descripción", "Especificación técnica", "Und", "P1", "P2", "P3", "P4", "P5",
           "Cantidad total", "Fuente", "Observaciones"],
    [7, 38, 58, 6, 9, 9, 9, 9, 9, 13, 7, 46])
C.freeze_panes = "E5"

DRS = "'Datos Revit'!"
FOS = "'Foso ascensor'!"
r = 4
chap_rows = []


def chapter(num, name):
    global r
    r += 1
    for col in range(1, 13):
        c = C.cell(row=r, column=col)
        c.fill, c.border = CHFILL, BRD
        c.font = BOLD
    C[f"A{r}"] = num
    C[f"B{r}"] = name
    chap_rows.append(r)


items = []  # (row, chapter_num)


def item(code, desc, spec, und, per=None, total=None, src="R", obs="", fmt=N2):
    """per: list of 5 formulas/values for P1..P5 (or None). total: formula if no per-floor."""
    global r
    r += 1
    put(C, f"A{r}", code); put(C, f"B{r}", desc); put(C, f"C{r}", spec); put(C, f"D{r}", und)
    if per:
        for i, v in enumerate(per):
            font = GREEN if isinstance(v, str) and "!" in v else BLACK
            put(C, f"{get_column_letter(5+i)}{r}", v, font, fmt)
        put(C, f"J{r}", f"=SUM(E{r}:I{r})", BOLD, fmt)
    else:
        for i in range(5):
            put(C, f"{get_column_letter(5+i)}{r}", None)
        font = GREEN if isinstance(total, str) and "!" in total else BLUE
        put(C, f"J{r}", total, Font(name=F, size=10, bold=True, color=font.color.rgb if font.color else None), fmt)
    put(C, f"K{r}", src); C[f"K{r}"].alignment = CENTER
    put(C, f"L{r}", obs)
    return r


def drv(row, colfirst="C"):
    """Per-floor links to a Datos Revit row starting at column colfirst."""
    start = ord(colfirst)
    return [f"={DRS}{chr(start+i)}{row}" for i in range(5)]


# 1 Preliminares
chapter("1", "PRELIMINARES")
item("1.1", "Localización, trazado y replanteo",
     "Con equipo de topografía (estación total / nivel). Incluye referencias de ejes y niveles, mojones y estacas.",
     "m²", total=f"={DRS}H{DR['area_cub']}+{DRS}H{DR['area_ext']}", src="R/D",
     obs="Área cubierta total P1–P5 + área exterior P1 (área de intervención por nivel).")
item("1.2", "Localización y replanteo del foso de ascensor",
     "Replanteo de ejes A-B / 1-2 según planta de geometría EST-02 (cotas a ejes: 1.10, 1.25, 1.19, 1.34 m).",
     "m²", total=f"={FOS}{G['TX']}*{FOS}{G['TY']}", src="E", obs="Huella 2.10 × 2.50 m.")

# 2 Foso
chapter("2", "FOSO DE ASCENSOR — CIMENTACIÓN Y ESTRUCTURA (EST-02)")
item("2.1", "Excavación manual en material común",
     "Hasta cota N-1.55 (fondo de solado), con sobreancho de 0.30 m por lado. Incluye perfilado de fondo. "
     "Precaución: zapatas vecinas en ejes A y B (NSZ-2.00 / NIZ-2.50) — no descalzar.",
     "m³", total=f"={FOS}B{FC['exc']}", src="E/D")
item("2.2", "Solado de limpieza e=0.05 m",
     "Concreto pobre f'c = 14 MPa (2.000 psi) bajo losa de fondo.", "m³",
     total=f"={FOS}B{FC['solado']}*(1+{SREF['desp_conc']})", src="D",
     obs="No dibujado en EST-02; recomendado. Incluye desperdicio.", fmt=N3)
item("2.3", "Concreto losa de fondo e=0.20 m",
     "Concreto f'c = 21 MPa (3.000 psi) con aditivo impermeabilizante integral. Cara superior a N-1.30, inferior N-1.50. "
     "Recubrimiento 75 mm contra suelo.", "m³",
     total=f"={FOS}B{FC['losa']}*(1+{SREF['desp_conc']})", src="E", obs="Incluye desperdicio. f'c NO indicado en EST-02: confirmar en notas generales (EST-01).", fmt=N3)
item("2.4", "Concreto muros de contención e=0.15 m",
     "Concreto f'c = 21 MPa (3.000 psi) con aditivo impermeabilizante integral, h = 1.20 m sobre losa (hasta N-0.10).",
     "m³", total=f"={FOS}B{FC['muro']}*(1+{SREF['desp_conc']})", src="E", obs="Incluye desperdicio.", fmt=N3)
item("2.5", "Concreto columnas 40×40 cm",
     "Concreto f'c = 21 MPa (3.000 psi). 4 und, de N-1.50 a N-0.10 (h = 1.40 m). Reciben perfil HE 200.",
     "m³", total=f"={FOS}B{FC['col']}*(1+{SREF['desp_conc']})", src="E", obs="Incluye desperdicio.", fmt=N3)
item("2.6", "Acero de refuerzo #4 (1/2\")",
     "Acero corrugado fy = 420 MPa (60.000 psi), NTC 2289. Losa doble parrilla #4 c/0.20; muros #4 c/0.20 en ambos sentidos; "
     "longitudinal de columnas 4 #4 por columna. Figurado según cartilla (hoja Foso ascensor).",
     "kg", total=f"={FOS}K{FC['k4']}*(1+{SREF['desp_acero']})", src="E/D", obs="Incluye desperdicio.")
item("2.7", "Acero de refuerzo #3 (3/8\") — flejes",
     "Acero corrugado fy = 420 MPa, NTC 2289. Flejes cerrados #3 @0.075 m, L = 1.48 m, gancho a 135°.",
     "kg", total=f"={FOS}K{FC['k3']}*(1+{SREF['desp_acero']})", src="E/D", obs="Incluye desperdicio.")
item("2.8", "Formaleta para muros y columnas del foso",
     "Formaleta metálica o en madera cepillada, con desmoldante. Ambas caras de muros + caras exteriores de columnas.",
     "m²", total=f"={FOS}B{FC['form']}", src="D")
item("2.9", "Impermeabilización interior del foso",
     "Mortero impermeable cementicio (2 manos cruzadas) en fondo y muros + mediacaña en encuentros muro-losa. "
     "Rematar a N-0.10.", "m²", total=f"={FOS}B{FC['imp']}", src="P",
     obs="No especificado en EST-02. Recomendado (foso bajo nivel de terreno).")
item("2.10", "Relleno compactado perimetral",
     "Material seleccionado de excavación o recebo, en capas ≤ 0.20 m compactadas al 95 % Proctor modificado.",
     "m³", total=f"={FOS}B{FC['rell']}", src="D")
item("2.11", "Retiro de sobrantes",
     "Cargue, transporte y disposición en sitio autorizado (escombrera).", "m³",
     total=f"={FOS}B{FC['ret']}", src="D", obs="Volumen suelto (con expansión).")
item("2.12", "Platina base A36 300×300×16 mm",
     "Acero ASTM A36, 4 perforaciones para pernos Ø5/8\". Nivelada con mortero sin retracción (grouting).",
     "un", total=f"={FOS}B{FC['pl_un']}", src="E", fmt=N0, obs="Peso total en hoja Foso ascensor.")
item("2.13", "Pernos de anclaje Ø5/8\" con gancho",
     "Barra roscada ASTM F1554 Gr. 36 (o A307), embebido 314 mm, tuerca y arandela. 4 por platina.",
     "un", total=f"={FOS}B{FC['pernos']}", src="E", fmt=N0)
item("2.14", "Soldadura de filete 5/16\"",
     "Electrodo E70XX, todo el contorno del perfil HE 200 a la platina.", "ml",
     total=f"={FOS}B{FC['sold']}", src="E/D", obs="Longitud aproximada por contorno de perfil.")
item("2.15", "Perfil HE 200 — columnas metálicas del ducto",
     "Perfil HE 200 (se asume HEB 200 = 61.3 kg/m), ASTM A572 Gr. 50 o A36, con anticorrosivo + esmalte.", "un",
     total=f"={FOS}B{FC['he']}", src="E/P", fmt=N0,
     obs="Cantidad en ml PENDIENTE: depende de la altura total del ducto (no está en EST-02).")

# 3 Mampostería
chapter("3", "MAMPOSTERÍA")
rm = item("3.1", "Muro en mampostería de arcilla",
          "Unidad de arcilla cocida NTC 4205 (tipo y espesor a confirmar), pega con mortero 1:4 tipo S (NTC 3329), "
          "juntas de 1 cm. Área neta medida en Revit (descuenta vanos, sin guardaescoba).",
          "m²", per=drv(DR["a_muros"]), src="R", obs="131 muros en P1–P5.")
item("3.2", "Unidades de mampostería (referencia)",
     "Bloque de arcilla 10×20×30 cm (supuesto). Incluye desperdicio.", "un",
     per=[f"={get_column_letter(5+i)}{rm}*{SREF['und_m2']}*(1+{SREF['desp_mamp']})" for i in range(5)],
     src="D", fmt=N0, obs="Ajustar rendimiento en Supuestos según la unidad especificada.")
item("3.3", "Mortero de pega 1:4",
     "Cemento gris + arena de pega, f'cm ≥ 12.5 MPa (tipo S).", "m³",
     per=[f"={get_column_letter(5+i)}{rm}*{SREF['mort_pega']}" for i in range(5)], src="D", fmt=N3)

# 4 Pañetes
chapter("4", "PAÑETES Y ACABADOS DE MUROS")
rp = item("4.1", "Pañete liso sobre muros (ambas caras)",
          "Mortero 1:4, e = 1.5 cm promedio, a plomo y regla. En fachada usar mortero 1:3 con impermeabilizante integral.",
          "m²", per=[f"={get_column_letter(5+i)}{rm}*{SREF['caras']}" for i in range(5)], src="D",
          obs="Bruto: descontar las zonas que lleven enchape (baños, cocinas) cuando se definan.")
item("4.2", "Estuco + pintura vinilo tipo 1 (3 manos)",
     "Estuco plástico + vinilo tipo 1 NTC 1335, color a definir. Fachada: pintura para exteriores (elastomérica o acrílica).",
     "m²", per=[f"={get_column_letter(5+i)}{rp}" for i in range(5)], src="D",
     obs="Misma área bruta que el pañete; descontar enchapes.")
item("4.3", "Dinteles sobre vanos (puertas y ventanas)",
     "Dintel en concreto 21 MPa o en perfil, apoyo 0.20 m por lado. Ancho del muro.", "ml",
     per=[f"=SUMPRODUCT(({DRS}$B${DR['puertas'][0]}:$B${DR['puertas'][-2]}+2*{SREF['apoyo_dintel']}),"
          f"{DRS}{get_column_letter(4+i)}${DR['puertas'][0]}:{get_column_letter(4+i)}${DR['puertas'][-2]})"
          f"+SUMPRODUCT(({DRS}$B${DR['vent'][0]}:$B${DR['vent'][-1]}+2*{SREF['apoyo_dintel']}),"
          f"{DRS}{get_column_letter(4+i)}${DR['vent'][0]}:{get_column_letter(4+i)}${DR['vent'][-1]})" for i in range(5)],
     src="D", obs="No incluye puertas de vestier Häfele (sin medida en el modelo).")

# 5 Pisos
chapter("5", "PISOS")
ra = item("5.1", "Alistado de piso (mortero de nivelación)",
          "Mortero 1:4, e = 4 cm, nivelado y afinado para recibir acabado.", "m²",
          per=drv(DR["area_cub"]), src="R", obs="Área cubierta por nivel (Revit).")
item("5.2", "Mortero para alistado", "Mortero 1:4.", "m³",
     per=[f"={get_column_letter(5+i)}{ra}*{SREF['e_alist']}" for i in range(5)], src="D", fmt=N3)
item("5.3", "Acabado de piso interior",
     "Tipo POR DEFINIR (porcelanato / cerámica antideslizante en baños y balcones). Incluye pegante y boquilla.",
     "m²", per=[f"={get_column_letter(5+i)}{ra}" for i in range(5)], src="R/P",
     obs="Área bruta; descontar muros si el área de Revit los incluye. Guardaescoba pendiente (perímetro de rooms).")
item("5.4", "Piso exterior / andén P1",
     "Concreto 21 MPa e = 8 cm sobre recebo compactado, o adoquín. Acabado antideslizante.", "m²",
     per=drv(DR["area_ext"]), src="R")

# 6 Cielorrasos
chapter("6", "CIELORRASOS")
item("6.1", "Cielorraso interior",
     "Tipo POR DEFINIR: pañete bajo placa + estuco y vinilo, o drywall en placa de yeso 1/2\". En baños usar placa RH.",
     "m²", per=[f"={get_column_letter(5+i)}{ra}" for i in range(5)], src="D/P",
     obs="Aproximado = área cubierta por nivel. Confirmar con plano de cielos (RCP).")

# 7 Puertas
chapter("7", "CARPINTERÍA — PUERTAS")
for k, prow in enumerate(DR["puertas"]):
    name, w, h, *_ = PUERTAS[k]
    med = f"{w:.3g} × {h:.3g} m" if w else "≈7.73 m² de hoja (total Revit)"
    if "garaje" in name:
        spec = "Puerta metálica enrollable o seccional, lámina cal. 20, con cerradura de seguridad."
    elif "principal" in name or "acceso" in name:
        spec = "Puerta de seguridad metálica o en madera maciza, marco metálico cal. 18, cerradura de seguridad y mirilla."
    elif "ventanal" in name or "social" in name:
        spec = "Puerta corrediza en aluminio con vidrio de seguridad templado 6 mm (NSR-10 Título K)."
    elif "Häfele" in name:
        spec = "Puertas corredizas de vestier con herraje Häfele, hoja en MDF/aglomerado enchapado."
    elif "corrediza" in name:
        spec = "Hoja corrediza entamborada en MDF, herraje de riel superior, manija embutida."
    elif "baño" in name:
        spec = "Hoja entamborada en MDF RH (resistente a humedad), marco en madera, cerradura de baño."
    else:
        spec = "Hoja entamborada en MDF, marco en madera, bisagras de 3½\" y cerradura de pomo."
    item(f"7.{k+1}", f"{name} — {med}", spec, "un", per=drv(prow, "D"), src="R", fmt=N0)

# 8 Ventanería
chapter("8", "VENTANERÍA Y VIDRIOS")
for k, vrow in enumerate(DR["vent"]):
    name, w, h, *_ = VENT[k]
    if h > 2:
        spec = "Perfilería en aluminio, vidrio de seguridad laminado o templado 6 mm (vano a piso, NSR-10 Título K)."
    else:
        spec = "Ventana en aluminio (proyectante o corrediza), vidrio incoloro 4 mm, empaques y sellos."
        if "baño" in name:
            spec = "Ventana en aluminio, vidrio esmerilado/sandblasting 4 mm (privacidad), con celosía de ventilación."
    item(f"8.{k+1}", f"{name} — {w:.3g} × {h:.3g} m", spec, "un", per=drv(vrow, "D"), src="R", fmt=N0)
n = len(DR["vent"])
item(f"8.{n+1}", "Área de vidrio en ventanas y ventanales", "Suma de ancho × alto × cantidad.", "m²",
     per=[f"=SUMPRODUCT({DRS}$B${DR['vent'][0]}:$B${DR['vent'][-1]},{DRS}$C${DR['vent'][0]}:$C${DR['vent'][-1]},"
          f"{DRS}{get_column_letter(4+i)}${DR['vent'][0]}:{get_column_letter(4+i)}${DR['vent'][-1]})" for i in range(5)],
     src="D", obs="No incluye vidrio de puertas-ventanal (ver ítem siguiente).")
pv = DR["puertas"][10]; ps = DR["puertas"][3]
item(f"8.{n+2}", "Área de vidrio de seguridad en puertas-ventanal y puerta de balcón", "Vidrio templado 6 mm.", "m²",
     per=[f"={DRS}$B${pv}*{DRS}$C${pv}*{DRS}{get_column_letter(4+i)}{pv}+{DRS}$B${ps}*{DRS}$C${ps}*{DRS}{get_column_letter(4+i)}{ps}" for i in range(5)],
     src="D")
item(f"8.{n+3}", "Alfajías en ventanas (excepto vanos a piso)",
     "Alfajía prefabricada en concreto o granito, con gotero, pendiente hacia el exterior.", "ml",
     per=[f"=SUMPRODUCT({DRS}$B${DR['vent'][0]}:$B${DR['vent'][1]},{DRS}{get_column_letter(4+i)}${DR['vent'][0]}:{get_column_letter(4+i)}${DR['vent'][1]})" for i in range(5)],
     src="D", obs="Solo ventanas de habitación y baño/cocina.")

# 9 Sanitarios
chapter("9", "APARATOS SANITARIOS, GRIFERÍAS E INCRUSTACIONES")
SPEC_SAN = {
    "san": ("Sanitario", "Sanitario de porcelana vitrificada, tanque de bajo consumo (≤ 4.8 L/descarga, NTC 920), con grifería de llenado y asiento."),
    "cab": ("Cabina de ducha", "Zona de ducha con piso en cerámica antideslizante, pendiente 1–2 % al sifón; cabina según modelo Revit."),
    "reg": ("Regadera de ducha", "Regadera + mezclador monocontrol cromado (agua fría/caliente), sistema ahorrador."),
    "dacc": ("Ducha accesible (con barras)", "Ducha a nivel sin sardinel, asiento abatible y barras de apoyo en acero inoxidable (NTC 6047 accesibilidad)."),
    "lav": ("Lavamanos de baño", "Lavamanos de porcelana (de incrustar u ovalín), con sifón, desagüe y grifería monocontrol ahorradora."),
    "mamp": ("Mampara de baño", "Mampara en vidrio templado 8 mm con perfilería en aluminio (NSR-10 Título K)."),
}
k = 0
for key in ["san", "cab", "reg", "dacc", "lav", "mamp"]:
    k += 1
    nm, sp = SPEC_SAN[key]
    item(f"9.{k}", nm, sp, "un", per=drv(DR["s_" + key], "D"), src="R", fmt=N0)
k += 1
item(f"9.{k}", "Grifería de lavamanos", "Monocontrol cromada, aireador ahorrador.", "un",
     per=drv(DR["s_lav"], "D"), src="D", fmt=N0, obs="1 por lavamanos.")
k += 1
item(f"9.{k}", "Mezclador de ducha", "Monocontrol empotrado, cartucho cerámico.", "un",
     per=[f"={DRS}{get_column_letter(4+i)}{DR['s_reg']}+{DRS}{get_column_letter(4+i)}{DR['s_dacc']}" for i in range(5)],
     src="D", fmt=N0, obs="Regaderas + ducha accesible.")
k += 1
item(f"9.{k}", "Juego de incrustaciones por baño",
     "Toallero, jabonera, portarrollos y gancho, en porcelana o acero inoxidable.", "jg",
     per=drv(DR["s_san"], "D"), src="D", fmt=N0, obs="1 juego por baño (= N° de sanitarios).")
k += 1
item(f"9.{k}", "Sifón de piso 3\" con rejilla", "PVC sanitario + rejilla en acero inoxidable.", "un",
     per=[f"={DRS}{get_column_letter(4+i)}{DR['s_cab']}+{DRS}{get_column_letter(4+i)}{DR['s_dacc']}" for i in range(5)],
     src="D", fmt=N0, obs="Mínimo 1 por ducha; agregar los de cocina/zona de ropas.")

# 10 Ascensor
chapter("10", "EQUIPO DE ASCENSOR")
item("10.1", "Ascensor de pasajeros",
     "Ascensor eléctrico sin cuarto de máquinas (MRL), capacidad y paradas POR DEFINIR (6 niveles), NTC 5926 / EN 81-20. "
     "Incluye cabina, puertas de piso, guías, contrapeso y control.", "gl", total=1, src="P", fmt=N0,
     obs="QA crítico: el ascensor no está modelado en Revit. El foso de EST-02 tiene luz libre 1.50 × 1.90 m: validar con el proveedor.")

# total rows: number formatting/borders done; set row heights
for row in range(5, r + 1):
    C.row_dimensions[row].height = None
C.auto_filter.ref = f"A4:L{r}"
LAST_C = r

# ---------------------------------------------------------------- RESUMEN
R = wb.create_sheet("Resumen", 0)
R["A1"] = "RESUMEN DE CANTIDADES PRINCIPALES — MULTIFAMILIAR HERRERA"
R["A1"].font = TITLE
R["A2"] = "Valores enlazados a las hojas de cálculo (se actualizan solos). Sin costos."
R["A2"].font = Font(name=F, size=9, italic=True)
hdr(R, 4, ["Indicador", "Cantidad", "Und", "Nota"], [52, 14, 8, 70])
RES = [
    ("Área cubierta P1–P5", f"={DRS}H{DR['area_cub']}", "m²", "Medida en Revit"),
    ("Área exterior (P1)", f"={DRS}H{DR['area_ext']}", "m²", ""),
    ("Muros de mampostería (área neta, una cara)", f"={DRS}H{DR['a_muros']}", "m²", "131 muros"),
    ("Pañete / estuco / pintura sobre muros (bruto, 2 caras)", f"=Cantidades!J{rp}", "m²", "Descontar enchapes"),
    ("Puertas (todas)", f"={DRS}I{DR['puertas_tot']}", "un", "Ver observación sobre KPI de 44"),
    ("Ventanas y ventanales", f"={DRS}I{DR['vent'][-1]+1}", "un", "32 convencionales + 6 paneles de muro cortina"),
    ("Aparatos sanitarios", f"={DRS}I{DR['s_mamp']+1}", "un", "Incluye mamparas"),
    ("Foso ascensor — concreto estructural (con desperdicio)", f"={FOS}B{FC['tot']}*(1+{SREF['desp_conc']})", "m³", "Losa + muros + columnas"),
    ("Foso ascensor — acero de refuerzo total (con desperdicio)", f"=({FOS}K{FC['k4']}+{FOS}K{FC['k3']})*(1+{SREF['desp_acero']})", "kg", "#4 + #3"),
    ("Foso ascensor — excavación", f"={FOS}B{FC['exc']}", "m³", ""),
]
rr = 4
for name, f, und, nota in RES:
    rr += 1
    put(R, f"A{rr}", name); put(R, f"B{rr}", f, GREEN, N2 if und != "un" else N0)
    put(R, f"C{rr}", und); put(R, f"D{rr}", nota)

# ---------------------------------------------------------------- OBSERVACIONES
O = wb.create_sheet("Observaciones")
O["A1"] = "OBSERVACIONES, INCONSISTENCIAS Y PENDIENTES"
O["A1"].font = TITLE
hdr(O, 3, ["N°", "Tipo", "Observación", "Acción recomendada"], [5, 16, 80, 60])
OBS = [
    ("Inconsistencia", "Área de muros: el aviso del dashboard dice 1,377.66 m² / 130 muros, pero la tabla por piso suma 1,359.15 m² / 131 muros.",
     "Se usó la tabla por piso (1,359.15 m²). Actualizar el aviso del dashboard."),
    ("Inconsistencia", "Puertas: el KPI dice 44 und, la tabla por función suma 42. En P2 la ficha del plano dice 8 puertas y la tabla 7.",
     "Re-extraer la tabla de puertas de Revit y conciliar (faltan 2 en el detalle)."),
    ("Inconsistencia", "Aparatos sanitarios: el KPI dice 61, el gráfico por nivel suma 56 y la tabla detallada suma 62.",
     "Se usó la tabla detallada (62). Verificar en Revit."),
    ("Faltante crítico", "El plano estructural recibido es EST-02 de 02: falta EST-01 (cimentación general, columnas, vigas, placas, escaleras).",
     "Enviar EST-01 para cuantificar concreto y acero del edificio completo."),
    ("Faltante crítico", "El ascensor no está modelado en Revit (QA crítico). El foso aparece en EST-02 entre ejes A-B / 1-2.",
     "Coordinar la ubicación del ducto en planta arquitectónica y la ficha del proveedor (foso útil 1.20 m, luz 1.50 × 1.90 m)."),
    ("Resuelto", "Columnas 40×40: confirmado 4 #4 por columna (16 barras en total). Se tomó L=1.85 m del detalle de refuerzo.",
     "Confirmado por el arquitecto. Verificar que la longitud de 1.85 m quede en la cartilla final."),
    ("Resuelto", "Flejes: confirmado fleje #3 @0.075 con L=1.48 m (80 flejes en total).",
     "Confirmado por el arquitecto."),
    ("EST-02", "Barras de losa: la planta indica L=2.10/2.50 y los cortes L=2.20/2.60. Se usaron las de los cortes (losa) y las de planta (horizontales de muros).",
     "Confirmar la cartilla de hierros con el calculista."),
    ("EST-02", "La lámina no indica f'c del concreto, fy del acero, recubrimientos generales ni tipo exacto de perfil (HEA/HEB 200). Se asumió 21 MPa, 420 MPa y HEB 200.",
     "Verificar en notas generales de EST-01."),
    ("EST-02", "El rótulo indica archivo digital '2026_06_24_TANQUE ALMACENAMIENTO_T2', parece una plantilla reutilizada.",
     "Solicitar corrección del rótulo."),
    ("Alcance", "Solo se cuantifican P1–P5. P6, P7 y la cubierta siguen en diseño.",
     "Actualizar cuando se terminen esos niveles en el modelo."),
    ("Pendiente", "Sin datos para: enchapes de baños y cocinas, guardaescobas, mesones, instalaciones hidrosanitarias, eléctricas y de gas, cubierta, escaleras y barandas, aparatos de cocina.",
     "Extraer de Revit las tablas de perímetro de rooms y de acabados, y los planos técnicos."),
    ("Supuesto", "Las cantidades derivadas (unidades de mampostería, mortero, pañetes, dinteles, cielos) usan los rendimientos de la hoja Supuestos.",
     "Ajustarlos según las especificaciones definitivas del proyecto."),
]
for i, (t, o, a) in enumerate(OBS, 1):
    put(O, f"A{3+i}", i); put(O, f"B{3+i}", t); put(O, f"C{3+i}", o); put(O, f"D{3+i}", a)

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != F:
                c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic,
                              color=c.font.color)
C.page_setup.orientation = "landscape"
C.page_setup.fitToWidth = 1
C.page_setup.fitToHeight = 0
C.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(OUT)
print("ok", LAST_C)
